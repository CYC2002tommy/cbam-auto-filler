"""Print the writable columns of given template rows with their header text.

    python tools/inspect_rows.py B_EmInst 17:14-16 98:95-97 113:110-112

Each argument is ROW:HEADER_FIRST-HEADER_LAST. Header text comes from the cached
(translated) values of the header rows directly above.
"""
import json
import sys
import warnings
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
MAP = json.loads((ROOT / "template" / "templateMap.json").read_text(encoding="utf-8"))
wb = openpyxl.load_workbook(ROOT / "template" / "cbam-template-v2.1.1.xlsx", data_only=True)

sheet = sys.argv[1]
ws = wb[sheet]
cells = MAP["cells"][sheet]
for spec in sys.argv[2:]:
    row_s, hdr = spec.split(":")
    row = int(row_s)
    h1, h2 = (int(x) for x in hdr.split("-"))
    cols = sorted((a for a in cells if a.rstrip("0123456789") and a[len(a.rstrip("0123456789")):] == str(row)),
                  key=lambda a: column_index_from_string(a.rstrip("0123456789")))
    print(f"\n{sheet} row {row}: {len(cols)} writable")
    for a in cols:
        col = a.rstrip("0123456789")
        ci = column_index_from_string(col)
        header = " / ".join(str(ws.cell(row=r, column=ci).value).strip() for r in range(h1, h2 + 1)
                            if ws.cell(row=r, column=ci).value not in (None, ""))
        rec = cells[a]
        lst = f" list={MAP['lists'][rec['l']][:4]}" if "l" in rec else ""
        print(f"  {col:>3} [{rec['t']}] {header[:90]}{lst}")
