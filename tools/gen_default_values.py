"""Generate data/defaultValues.json from the EU's default-values workbook.

The app looks up the EU default for Taiwan here instead of asking the user to type it.
Values are the raw totals from the regulation; the mark-up is applied in the app.
Re-run this script whenever the Commission publishes a new version of the workbook.
The workbook is not committed: download it from the TAXUD "Default values and benchmarks"
section (URL in the output's `source` field) and save it under the name in SRC.

    python tools/gen_default_values.py            # regenerate
    python tools/gen_default_values.py --report   # regenerate and print spot checks
"""
import hashlib
import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "reference" / "cbam-default-values-v2-20260806.xlsx"
OUT = ROOT / "data" / "defaultValues.json"
COUNTRY = "Taiwan"
# IR 2025/2621 Annex I: where a listed country shows "-" or no value, the value from
# "Other countries and territories" applies. The sheet name is truncated to 31 chars.
FALLBACK = "_Other Countries and Territorie"


def number(cell):
    text = str(cell).strip() if cell is not None else ""
    return float(text.replace(",", ".")) if re.fullmatch(r"\d+,\d+|\d+", text) else None


def rows(ws):
    """(code, description, sector, total) for every good that has a code, in sheet order."""
    sector = None
    for code, desc, _direct, _indirect, total, *_ in ws.iter_rows(min_row=3, values_only=True):
        if code and not any([desc, _direct, _indirect, total]):
            sector = str(code).strip()
            continue
        if code and str(total).strip() != "see below":
            yield " ".join(str(code).split()), str(desc).strip(), sector, number(total)


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    history = [r for r in wb["Version History"].iter_rows(min_row=3, values_only=True) if r[0]]
    version, date, notes = history[-1][:3]

    fallback = {code: total for code, _, _, total in rows(wb[FALLBACK])}
    goods = []
    for code, desc, sector, total in rows(wb[COUNTRY]):
        from_other = total is None
        value = fallback.get(code) if from_other else total
        if value is None:
            continue  # no value in either table: the regulation sets none for this good
        goods.append({"cn": code, "description": desc, "sector": sector,
                      "total": value, "fromOther": from_other})

    out = {
        "country": COUNTRY,
        "version": int(version),
        "date": date.strftime("%Y-%m-%d"),
        "notes": notes,
        "legalBasis": "IR 2025/2621 Annex I, as corrected by IR 2026/1740",
        "source": "https://taxation-customs.ec.europa.eu/carbon-border-adjustment-mechanism/cbam-legislation-and-guidance_en#default-values-and-benchmarks",
        "sha256": hashlib.sha256(SRC.read_bytes()).hexdigest(),
        "goods": goods,
    }
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")
    print(f"{OUT.relative_to(ROOT)}: {len(goods)} goods, {sum(g['fromOther'] for g in goods)} from 'Other countries', version {out['version']} ({out['date']})")

    if "--report" in sys.argv:
        by = {g["cn"]: g for g in goods}
        for cn in ("7318 15", "7318 16", "7601", "2523 29 00", "3102 10 19", "2804 10 00", "7218 10 00"):
            g = by.get(cn)
            print(f"  {cn:12} {g['total'] if g else '—':>7}  {g['sector'] if g else ''}{'  (from Other countries)' if g and g['fromOther'] else ''}")


if __name__ == "__main__":
    main()
